"""Offline contract checks; optional isolated PostgreSQL DDL/query check (always rolled back)."""
import argparse
import json
from pathlib import Path
import sys
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def offline():
    from openpyxl import load_workbook
    from fastapi import HTTPException
    from src.gateway.forms.validation import value_for
    catalog = {d["code"]: d for d in json.loads((ROOT / "src/gateway/forms/catalog.json").read_text(encoding="utf8"))}
    book = load_workbook(ROOT / "tests/fixtures/forms/synthetic-valid.xlsx", data_only=False)
    rows = list(book["GIA_LAP"].values)
    assert len(rows) == 3 and rows[1][0] == "SYNTHETIC_001"
    for row in rows[1:]:
        for code, value in zip(("tuoi", "kho_tho_khi_nam", "can_nang"), row[2:5]):
            assert value_for(catalog[code], value) == value
    assert rows[1][2] == 0 and rows[1][3] is False and rows[1][4] is None
    book.close()
    invalid = load_workbook(ROOT / "tests/fixtures/forms/synthetic-invalid.xlsx", data_only=False)
    invalid_rows = list(invalid["GIA_LAP"].values)
    try:
        value_for(catalog["tuoi"], invalid_rows[1][2])
        raise AssertionError("invalid type accepted")
    except HTTPException:
        pass
    assert invalid_rows[2][0] is None and invalid_rows[2][1] is None
    invalid.close()
    for value in (None, [], [{"name": "SYNTHETIC_ANALYTE", "value": 0, "unit": None, "observed_at": None}]):
        assert value_for(catalog["dien_giai"], value) == value
    assert len(catalog) == 206 and sum(d["type"] != "unavailable" for d in catalog.values()) == 72
    print(json.dumps({"offline": "passed", "catalog_rows": len(catalog), "valid_excel_rows": 2, "invalid_type_rejected": 1, "missing_identity_fixture_rows": 1}))


def database(url):
    import psycopg
    from psycopg.types.json import Jsonb
    conn = psycopg.connect(url)
    try:
        conn.execute("SET LOCAL ROLE hf_demo_owner")
        conn.execute("CREATE SCHEMA s1_form_qa_20261023")
        conn.execute("SET LOCAL search_path TO s1_form_qa_20261023")
        conn.execute("CREATE TABLE app_user(id uuid PRIMARY KEY)")
        conn.execute("CREATE TABLE patient_case(id uuid PRIMARY KEY)")
        for name in ("002_dynamic_forms.sql", "003_form_imports.sql"):
            sql = (ROOT / "database/migrations" / name).read_text(encoding="utf8")
            sql = sql.replace("BEGIN;", "").replace("COMMIT;", "")
            conn.execute(sql)
            conn.execute(sql)  # additive migration replay
        actor, patient, visit, mutation = (uuid4() for _ in range(4))
        conn.execute("INSERT INTO app_user VALUES(%s)", (actor,))
        conn.execute("INSERT INTO form_patient(id,code,owner_id) VALUES(%s,'SYNTHETIC_001',%s)", (patient, actor))
        values = {"zero": 0, "false": False, "unknown": None, "empty": [], "repeated": [{"name": "SYNTHETIC", "value": 0}]}
        conn.execute("INSERT INTO form_visit(id,patient_id,external_id,created_by,values_json) VALUES(%s,%s,'SYNTHETIC_VISIT_001',%s,%s)", (visit, patient, actor, Jsonb(values)))
        conn.execute("INSERT INTO form_history(visit_id,revision,mutation_id,request_hash,actor_id,values_json,sources_json,definitions_json,catalog_revision) VALUES(%s,1,%s,'synthetic',%s,%s,'{}','{}',1)", (visit, mutation, actor, Jsonb(values)))
        args = (uuid4(), actor)
        conn.execute("INSERT INTO form_import_batch(id,actor_id,source_kind,content_hash,mapping,result) VALUES(%s,%s,'excel','synthetic','{}','{}')", args)
        rejected = 0
        for statement, params in [
            ("INSERT INTO form_import_batch(id,actor_id,source_kind,content_hash,mapping,result) VALUES(%s,%s,'excel','synthetic','{}','{}')", (uuid4(), actor)),
            ("INSERT INTO form_visit(id,patient_id,created_by) VALUES(%s,%s,%s)", (uuid4(), uuid4(), actor))]:
            try:
                with conn.transaction():
                    conn.execute(statement, params)
            except psycopg.IntegrityError:
                rejected += 1
        actual = conn.execute("SELECT values_json FROM form_visit WHERE id=%s", (visit,)).fetchone()[0]
        assert actual == values
        assert conn.execute("SELECT count(*) FROM form_import_batch").fetchone()[0] == 1
        assert conn.execute("SELECT count(*) FROM form_history").fetchone()[0] == 1
        plan = conn.execute("EXPLAIN (ANALYZE, FORMAT JSON) SELECT id,revision,values_json FROM form_visit WHERE patient_id=%s", (patient,)).fetchone()[0]
        assert rejected == 2
        print(json.dumps({"postgres": "passed", "migrations_replayed": 2, "patients": 1, "visits": 1, "history": 1, "imports": 1, "rejected_constraints": rejected, "query_plan": plan[0]["Plan"]["Node Type"], "execution_ms": plan[0]["Execution Time"], "cleanup": "transaction rollback"}))
    finally:
        conn.rollback()
        conn.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--database-url")
    args = parser.parse_args()
    offline()
    if args.database_url:
        database(args.database_url)
