"""Bounded Excel parser, read-only signed previews and atomic imports."""
from __future__ import annotations
import base64
from datetime import date, datetime, timezone
import hashlib
import hmac
from io import BytesIO
import json
import re
import secrets
import time
from typing import Any, Literal
from uuid import UUID, uuid4
from zipfile import ZipFile
from xml.etree.ElementTree import iterparse

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pydantic import Field
from psycopg.types.json import Jsonb

from ..database import staff_connection
from .router import catalog, fingerprint, form_actor, persist, visit_access
from .validation import Input, value_for

router = APIRouter(prefix='/api/forms/imports', tags=['Nhập biểu mẫu'])
LIMITS = {'records': 100, 'fields': 500,
          'array_items': 200, 'excel_bytes': 4194304, 'excel_expanded_bytes': 33554432,
          'excel_rows': 1000, 'excel_columns': 500, 'preview_seconds': 1800}


def problem(path, reason):
    raise HTTPException(422, detail={'path': path, 'reason': reason})


class ExcelFile(Input):
    filename: str = Field(min_length=1, max_length=200)
    content_base64: str = Field(max_length=5592408)


def workbook(body: ExcelFile):
    if not body.filename.lower().endswith('.xlsx'):
        problem('/filename', 'Chỉ hỗ trợ .xlsx; không đổi đuôi .xls để nhập')
    try:
        raw = base64.b64decode(body.content_base64, validate=True)
        if len(raw) > LIMITS['excel_bytes']:
            problem('/file', 'Excel vượt 4 MiB')
        with ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(m.file_size for m in members) > LIMITS['excel_expanded_bytes']:
                problem('/file', 'Workbook vượt giới hạn giải nén 32 MiB hoặc 1000 mục')
            if any(m.filename.lower().endswith('vbaproject.bin') for m in members):
                problem('/file', 'Không nhận macro')
            # Bound stored coordinates before read-only openpyxl can expand a
            # dimensionless sheet's sparse rows into an enormous iterator.
            for member in members:
                if member.filename.startswith('xl/worksheets/') and member.filename.endswith('.xml'):
                    with archive.open(member) as source:
                        for _, element in iterparse(source, events=('end',)):
                            tag = element.tag.rsplit('}', 1)[-1]
                            if tag == 'row' and int(element.get('r', '0')) > 1100:
                                problem('/sheet', 'Sheet vượt 1100 hàng gồm tiêu đề')
                            if tag == 'c':
                                coordinate = element.get('r', '')
                                match = re.fullmatch(r'([A-Z]+)(\d+)', coordinate)
                                if match:
                                    column = 0
                                    for char in match[1]:
                                        column = column * 26 + ord(char) - 64
                                    if column > 500 or int(match[2]) > 1100:
                                        problem('/sheet', 'Sheet vượt 500 cột hoặc 1100 hàng gồm tiêu đề')
                            element.clear()
        return load_workbook(BytesIO(raw), read_only=True, data_only=False)
    except HTTPException:
        raise
    except Exception:
        problem('/file', 'Không đọc được workbook .xlsx hợp lệ')


class ExcelInspect(ExcelFile):
    sheet: str | None = Field(default=None, max_length=200)
    header_row: int = Field(default=1, ge=1, le=100)


def sheet_for(book, name):
    if name not in book.sheetnames:
        problem('/sheet', 'Chọn sheet có trong workbook')
    sheet = book[name]
    if sheet.max_column is None or sheet.max_row is None:
        try:
            sheet.calculate_dimension(force=True)
        except (ValueError, UnboundLocalError):
            # Empty dimensionless XML is a valid empty sheet.
            sheet._max_row, sheet._max_column = 1, 1
    if (sheet.max_column or 0) > LIMITS['excel_columns'] or (sheet.max_row or 0) > LIMITS['excel_rows'] + 100:
        problem('/sheet', 'Sheet vượt 500 cột hoặc 1000 dòng dữ liệu + 100 dòng tiêu đề')
    return sheet


def cell_display(cell):
    if cell.value is None:
        return None
    if isinstance(cell.value, (date, datetime)):
        return cell.value.isoformat()
    return cell.value


@router.get('/limits')
def limits(actor=Depends(form_actor)):
    return LIMITS


@router.post('/excel/inspect')
def excel_inspect(body: ExcelInspect, actor=Depends(form_actor)):
    book = workbook(body)
    try:
        result = {'sheets': book.sheetnames, 'limits': LIMITS, 'business_writes': False}
        if body.sheet is not None:
            sheet = sheet_for(book, body.sheet)
            if body.header_row > (sheet.max_row or 0):
                problem('/header_row', 'Dòng tiêu đề không có trong sheet')
            rows = list(sheet.iter_rows(min_row=body.header_row, max_row=min(sheet.max_row, body.header_row + 5)))
            result.update(columns=[{'column': i + 1, 'header': str(c.value or ''), 'letter': get_column_letter(i + 1)} for i, c in enumerate(rows[0])] if rows else [],
                          sample=[{'row': body.header_row + i + 1, 'values': [cell_display(c) for c in cells]} for i, cells in enumerate(rows[1:])])
        return result
    finally:
        book.close()


class ColumnMap(Input):
    column: int = Field(ge=1, le=500)
    target: str = Field(min_length=1, max_length=80)
    unit_confirmed: bool = False
    conversion: Literal['native', 'dot_number', 'boolean_01', 'boolean_true_false', 'iso_date'] = 'native'


class ExplicitTarget(Input):
    row: int = Field(ge=2, le=1100)
    visit_id: UUID
    confirmed: Literal[True]


class ExcelPreview(ExcelFile):
    sheet: str = Field(min_length=1, max_length=200)
    header_row: int = Field(ge=1, le=100)
    mapping: list[ColumnMap] = Field(min_length=1, max_length=500)
    patient_column: int | None = Field(default=None, ge=1, le=500)
    visit_column: int | None = Field(default=None, ge=1, le=500)
    explicit_target: ExplicitTarget | None = None


def excel_value(cell, field, mapping):
    value = cell.value
    if cell.data_type == 'f':
        raise ValueError('Ô công thức không được thực thi; cần giá trị nguồn đã xác minh')
    if cell.data_type == 'e':
        raise ValueError('Ô Excel chứa lỗi')
    if value is None or value == '':
        return None
    if field.get('unit') and not mapping.unit_confirmed:
        raise ValueError('Cần xác nhận đơn vị nguồn trùng đơn vị đích; không tự chuyển đổi')
    if mapping.conversion == 'dot_number' and isinstance(value, str):
        if not re.fullmatch(r'-?(?:\d+(?:\.\d+)?|\.\d+)', value.strip()):
            raise ValueError('Cần số thập phân dấu chấm, không dấu phân cách nghìn')
        value = float(value) if '.' in value else int(value)
    elif mapping.conversion in {'boolean_01', 'boolean_true_false'}:
        codes = {'0': False, '1': True} if mapping.conversion == 'boolean_01' else {'true': True, 'false': False}
        key = str(value).lower()
        if key not in codes:
            raise ValueError('Không khớp bảng mã Có/Không đã chọn')
        value = codes[key]
    elif mapping.conversion == 'iso_date' and isinstance(value, (date, datetime)):
        if field['type'] != 'date':
            raise ValueError('Ngày Excel thiếu múi giờ; không tự gán cho thời điểm')
        value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    elif isinstance(value, (date, datetime)):
        raise ValueError('Chọn chuyển ngày ISO rõ ràng; thời điểm cần múi giờ')
    return value


def _secret(request):
    return request.app.state.form_preview_secret


def sign_plan(request, plan):
    payload = base64.urlsafe_b64encode(json.dumps(plan, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=str).encode()).decode()
    signature = hmac.new(_secret(request), payload.encode(), hashlib.sha256).hexdigest()
    return payload + '.' + signature


def read_plan(request, token, actor):
    try:
        payload, signature = token.rsplit('.', 1)
        expected = hmac.new(_secret(request), payload.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(expected, signature):
            raise ValueError()
        plan = json.loads(base64.urlsafe_b64decode(payload))
        if plan['actor'] != str(actor.user_id) or plan['expires'] < time.time():
            raise ValueError()
        return plan
    except Exception:
        raise HTTPException(409, 'Bản xem trước hết hạn/không hợp lệ hoặc máy chủ đã khởi động lại; xem trước lại') from None


def preview_plan(conn, actor, records, revision, definitions, kind, mapping, source_name):
    rows = []
    identities = [('visit_uuid', r['_explicit_visit_id']) if r.get('_explicit_visit_id') else (r.get('patient_code'), r.get('visit_external_id')) for r in records]
    for index, record in enumerate(records):
        errors = list(record.get('_errors', []))
        path = f'/records/{index}'
        identity = (record.get('patient_code'), record.get('visit_external_id'))
        row = {'index': index, 'patient_code': identity[0], 'visit_external_id': identity[1], 'errors': errors, 'fields': [], 'values': {}, 'sources': {}}
        target = conn.execute('''SELECT v.* FROM form_visit v JOIN form_patient p ON p.id=v.patient_id
            WHERE p.code=%s AND v.external_id=%s AND (p.owner_id=%s OR EXISTS
            (SELECT 1 FROM form_access a WHERE a.patient_id=p.id AND a.user_id=%s))''',
            (*identity, actor.user_id, actor.user_id)).fetchall() if all(identity) else []
        if record.get('_explicit_visit_id'):
            current = visit_access(conn, actor, UUID(record['_explicit_visit_id']))
            target = [current]
        if identities.count(identities[index]) > 1:
            errors.append({'path': path, 'reason': 'Nhiều bản ghi cùng hồ sơ/lượt; chọn rõ một bản ghi, không lấy dòng cuối'})
        if len(target) != 1:
            errors.append({'path': path, 'reason': 'Không khớp chính xác hồ sơ/lượt trong phạm vi tài khoản; cần đối chiếu'})
        else:
            row.update(visit_id=str(target[0]['id']), expected_revision=target[0]['revision'])
        mapped_targets = set()
        for source_code, raw in record['values'].items():
            target_code = source_code
            location = record.get('_locations', {}).get(source_code, path + '/values/' + source_code.replace('~', '~0').replace('/', '~1'))
            if target_code is None:
                row['fields'].append({'source': source_code, 'path': location, 'status': 'unmapped', 'raw': raw})
                continue
            item = {'source': source_code, 'code': target_code, 'path': location, 'raw': record.get('_raw', {}).get(source_code, raw)}
            field = definitions.get(target_code)
            try:
                if not field or not field['active'] or field['type'] == 'unavailable':
                    raise ValueError('Mã trường lạ, chưa đặc tả hoặc đã ngừng dùng')
                if target_code in mapped_targets:
                    raise ValueError('Hai nguồn cùng ánh xạ một trường; không tự ghi đè')
                mapped_targets.add(target_code)
                value = value_for(field, raw)
                current = target[0]['values_json'].get(target_code) if len(target) == 1 else None
                status = 'source_null' if value is None else 'unchanged' if value == current else 'fill' if current is None else 'conflict'
                item.update(label=field['label'], incoming=value, current=current, status=status, current_source=target[0]['sources_json'].get(target_code) if len(target) == 1 else None)
                row['values'][target_code] = value
                row['sources'][target_code] = {'kind': kind, 'location': location, 'source_name': source_name, 'source_key': source_code}
            except (ValueError, HTTPException) as exc:
                reason = str(exc) if isinstance(exc, ValueError) else str(exc.detail)
                item.update(status='error', reason=reason)
                errors.append({'path': location, 'field': target_code, 'reason': reason})
            row['fields'].append(item)
        row['status'] = 'error' if errors else 'eligible'
        rows.append(row)
    return {'actor': str(actor.user_id), 'expires': int(time.time()) + LIMITS['preview_seconds'], 'catalog_revision': revision,
            'source_kind': kind, 'source_name': source_name, 'mapping': mapping, 'mapping_version': '1.0', 'records': rows}


def response_for(request, plan):
    if len(json.dumps(plan, ensure_ascii=False, default=str).encode()) > 8 * 1024 * 1024:
        problem('/preview', 'Kế hoạch vượt 8 MiB; chọn ít dòng/cột hơn để nhập')
    return {'preview_token': sign_plan(request, plan), 'catalog_revision': plan['catalog_revision'], 'records': plan['records'],
            'eligible': sum(r['status'] == 'eligible' for r in plan['records']), 'errors': sum(r['status'] == 'error' for r in plan['records']),
            'business_writes': False, 'limits': LIMITS, 'mapping_version': '1.0'}


@router.post('/excel/preview')
def excel_preview(body: ExcelPreview, request: Request, actor=Depends(form_actor)):
    book = workbook(body)
    try:
        sheet = sheet_for(book, body.sheet)
        if len({m.column for m in body.mapping}) != len(body.mapping) or len({m.target for m in body.mapping}) != len(body.mapping):
            problem('/mapping', 'Mỗi cột và trường đích chỉ ánh xạ một lần')
        if body.explicit_target is None and (body.patient_column is None or body.visit_column is None):
            problem('/identity', 'Chọn cột mã hồ sơ và mã lượt hoặc đối chiếu một dòng với lượt đích rõ ràng')
        records = []
        with staff_connection(request, actor) as conn:
            revision, definitions = catalog(conn)
            explicit = None
            if body.explicit_target:
                v = visit_access(conn, actor, body.explicit_target.visit_id)
                p = conn.execute('SELECT code FROM form_patient WHERE id=%s', (v['patient_id'],)).fetchone()
                explicit = (p['code'], v['external_id'], str(v['id']))
                if body.explicit_target.row <= body.header_row or body.explicit_target.row > sheet.max_row:
                    problem('/explicit_target/row', 'Dòng phải sau tiêu đề và có trong sheet')
            for mapping in body.mapping:
                if mapping.target not in definitions:
                    problem('/mapping', 'Trường đích không tồn tại: ' + mapping.target)
            for row_number, cells in enumerate(sheet.iter_rows(min_row=body.header_row + 1, max_row=sheet.max_row), start=body.header_row + 1):
                if body.explicit_target and row_number != body.explicit_target.row:
                    continue
                if all(c.value is None for c in cells):
                    continue
                if len(records) >= LIMITS['records']:
                    problem('/records', 'Tối đa 100 dòng không trống mỗi lần; chọn một dòng để đối chiếu nếu cần')
                record = {'values': {}, '_errors': [], '_locations': {}, '_raw': {}}
                def cell_at(column):
                    if column is None or column > len(cells):
                        raise ValueError('Cột đã chọn không có trong sheet')
                    return cells[column - 1]
                if explicit:
                    record.update(patient_code=explicit[0], visit_external_id=explicit[1], _explicit_visit_id=explicit[2])
                else:
                    for prop, column in [('patient_code', body.patient_column), ('visit_external_id', body.visit_column)]:
                        try:
                            cell = cell_at(column)
                            if cell.data_type == 'f' or not isinstance(cell.value, str) or not cell.value:
                                raise ValueError('Định danh phải là chuỗi có sẵn; không đoán số 0 đầu hoặc đọc công thức')
                            record[prop] = cell.value
                        except ValueError as exc:
                            record[prop] = None
                            record['_errors'].append({'path': f'{body.sheet}!row{row_number}/column{column}', 'reason': str(exc)})
                for mapping in body.mapping:
                    try:
                        cell = cell_at(mapping.column)
                        record['_locations'][mapping.target] = f'{body.sheet}!{get_column_letter(mapping.column)}{row_number}'
                        record['_raw'][mapping.target] = cell_display(cell)
                        if definitions[mapping.target]['type'] in {'medications', 'results', 'multiselect'}:
                            raise ValueError('Cấu trúc danh sách Excel chưa được xác minh; bổ sung bằng biểu mẫu có cấu trúc, không gộp hoặc lấy dòng cuối')
                        record['values'][mapping.target] = excel_value(cell, definitions[mapping.target], mapping)
                    except (ValueError, HTTPException) as exc:
                        record['_errors'].append({'path': f'{body.sheet}!row{row_number}/column{mapping.column}', 'field': mapping.target,
                                                  'raw': record['_raw'].get(mapping.target),
                                                  'reason': str(exc) if isinstance(exc, ValueError) else str(exc.detail)})
                records.append(record)
            if not records:
                problem('/sheet', 'Không có dòng dữ liệu được chọn')
            metadata = {'columns': [m.model_dump() for m in body.mapping], 'sheet': body.sheet, 'header_row': body.header_row,
                        'patient_column': body.patient_column, 'visit_column': body.visit_column,
                        'explicit_target': body.explicit_target.model_dump(mode='json') if body.explicit_target else None}
            plan = preview_plan(conn, actor, records, revision, definitions, 'excel', metadata, body.filename)
        return response_for(request, plan)
    finally:
        book.close()


class Decision(Input):
    index: int = Field(ge=0, le=99)
    override_fields: list[str] = Field(default_factory=list, max_length=500)


class Commit(Input):
    preview_token: str = Field(min_length=10, max_length=12000000)
    records: list[Decision] = Field(min_length=1, max_length=100)


@router.post('/commit')
def commit(body: Commit, request: Request, actor=Depends(form_actor)):
    plan = read_plan(request, body.preview_token, actor)
    if len({d.index for d in body.records}) != len(body.records):
        problem('/records', 'Không chọn lặp bản ghi')
    selected = []
    for decision in body.records:
        if decision.index >= len(plan['records']):
            problem('/records', 'Bản ghi không có trong preview')
        row = plan['records'][decision.index]
        if row['status'] != 'eligible':
            problem('/records/' + str(decision.index), 'Bản ghi còn lỗi, không được ghi')
        conflicts = {f['code'] for f in row['fields'] if f['status'] == 'conflict'}
        if len(set(decision.override_fields)) != len(decision.override_fields) or set(decision.override_fields) - conflicts:
            problem('/records/' + str(decision.index), 'Chỉ chọn ghi đè cụ thể các trường xung đột')
        selected.append((row, set(decision.override_fields)))
    content_hash = fingerprint(sorted([{'patient_code': r['patient_code'], 'visit_external_id': r['visit_external_id'], 'visit_id': r['visit_id'], 'values': r['values'], 'overrides': sorted(o)} for r, o in selected], key=lambda r: r['visit_id']))
    with staff_connection(request, actor) as conn:
        # Serialize the identical retry before revisions; a lost response can
        # replay the durable result after a newer preview without new writes.
        lock_key = int.from_bytes(hashlib.sha256((str(actor.user_id) + content_hash).encode()).digest()[:8], 'big', signed=True)
        conn.execute('SELECT pg_advisory_xact_lock(%s)', (lock_key,))
        previous = conn.execute('SELECT result FROM form_import_batch WHERE actor_id=%s AND content_hash=%s', (actor.user_id, content_hash)).fetchone()
        if previous:
            for row, _ in selected:
                visit_access(conn, actor, UUID(row['visit_id']))
            return {**previous['result'], 'replayed': True}
        revision, definitions = catalog(conn)
        if revision != plan['catalog_revision']:
            raise HTTPException(409, 'Danh mục đã thay đổi; xem trước lại')
        locked = {}
        for row, _ in sorted(selected, key=lambda pair: pair[0]['visit_id']):
            current = visit_access(conn, actor, UUID(row['visit_id']), lock=True)
            if current['revision'] != row['expected_revision']:
                raise HTTPException(409, 'Dữ liệu lượt đã thay đổi; xem trước và đối chiếu lại')
            locked[row['visit_id']] = current
        batch_id = uuid4()
        outcomes = []
        for row, overrides in selected:
            current = locked[row['visit_id']]
            changes, sources = {}, {}
            skipped = []
            for item in row['fields']:
                status, code = item['status'], item.get('code')
                if status == 'fill' or status == 'conflict' and code in overrides:
                    changes[code] = value_for(definitions[code], row['values'][code])
                    sources[code] = {**row['sources'][code], 'actor_id': str(actor.user_id), 'recorded_at': datetime.now(timezone.utc).isoformat(),
                                     'batch_id': str(batch_id), 'mapping_version': plan['mapping_version']}
                else:
                    skipped.append({'field': code or item.get('source'), 'reason': status})
            new_revision = current['revision']
            if changes:
                saved = persist(conn, actor, current, revision, definitions, changes, sources, uuid4(), content_hash)
                new_revision = saved['revision']
            outcomes.append({'index': row['index'], 'patient_code': row['patient_code'], 'visit_external_id': row['visit_external_id'],
                             'visit_id': row['visit_id'], 'written_fields': len(changes), 'skipped': skipped, 'revision': new_revision})
        result = {'batch_id': str(batch_id), 'written_records': sum(o['written_fields'] > 0 for o in outcomes),
                  'written_fields': sum(o['written_fields'] for o in outcomes), 'skipped_fields': sum(len(o['skipped']) for o in outcomes),
                  'records': outcomes, 'unselected_records': len(plan['records']) - len(selected),
                  'unselected': [{'index': r['index'], 'status': r['status'], 'errors': r['errors']} for r in plan['records'] if r['index'] not in {d.index for d in body.records}],
                  'replayed': False,
                  'source_kind': plan['source_kind'], 'source_name': plan['source_name'], 'mapping_version': plan['mapping_version']}
        conn.execute('''INSERT INTO form_import_batch(id,actor_id,source_kind,content_hash,mapping,mapping_version,result)
            VALUES(%s,%s,%s,%s,%s,%s,%s)''', (batch_id, actor.user_id, plan['source_kind'], content_hash, Jsonb(plan['mapping']), plan['mapping_version'], Jsonb(result)))
    return result


@router.get('/history')
def import_history(request: Request, actor=Depends(form_actor)):
    with staff_connection(request, actor) as conn:
        return conn.execute('SELECT id,source_kind,mapping,mapping_version,result,created_at FROM form_import_batch WHERE actor_id=%s ORDER BY created_at DESC LIMIT 100', (actor.user_id,)).fetchall()


@router.get('/template.xlsx')
def excel_template(actor=Depends(form_actor)):
    book = Workbook()
    sheet = book.active
    sheet.title = 'GIA_LAP'
    sheet.append(['Mã hồ sơ', 'Mã lượt', 'Cân nặng (kg)', 'Khó thở khi nằm (true/false)', 'Ghi chú giả lập'])
    sheet.append(['SYN-FORM-001', 'VISIT-001-1', 61, 'false', 'MẪU GIẢ LẬP — không phải file HIS'])
    sheet.append(['SYN-FORM-002', 'VISIT-002-1', None, 'true', 'MẪU GIẢ LẬP — ô trống không xóa đích'])
    output = BytesIO()
    book.save(output)
    book.close()
    return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="synthetic-form-template.xlsx"'})
